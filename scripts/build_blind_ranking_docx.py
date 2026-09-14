#!/usr/bin/env python3
"""Build a source-blinded, four-way narrative ranking questionnaire."""

from __future__ import annotations

import argparse
import hashlib
import json
import random
from collections import defaultdict
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_WORK_DIR = ROOT / "data/human_pilot/private/api_family_pilot20"
DEFAULT_OUTPUT = ROOT / "deliverables/中文叙事质量人工盲评排序问卷_20题.docx"
DEFAULT_CODEBOOK = DEFAULT_WORK_DIR / "human_ranking_instrument_codebook.json"
FONT = "Noto Sans CJK SC"
INK = "243447"
BLUE = "245B78"
MUTED = "687684"
LIGHT = "EEF3F6"
BORDER = "BAC7D0"
TOTAL_WIDTH_DXA = 9360


def read_jsonl(path: Path) -> list[dict[str, Any]]:
    with path.open("r", encoding="utf-8") as handle:
        return [json.loads(line) for line in handle if line.strip()]


def set_cell_margins(cell: Any, top: int = 100, start: int = 120, bottom: int = 100, end: int = 120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table: Any, widths: list[int], indent: int = 120) -> None:
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def set_row_cant_split(row: Any) -> None:
    """Prevent a table row from being divided between two pages."""
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:cantSplit")) is None:
        tr_pr.append(OxmlElement("w:cantSplit"))


def set_repeat_table_header(row: Any) -> None:
    """Repeat a table's header row when the table continues on another page."""
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:tblHeader")) is None:
        tr_pr.append(OxmlElement("w:tblHeader"))


def keep_table_rows_together(table: Any) -> None:
    """Keep a short table together by chaining all cell paragraphs."""
    paragraphs = [p for row in table.rows for cell in row.cells for p in cell.paragraphs]
    for p in paragraphs[:-1]:
        p.paragraph_format.keep_with_next = True
    for row in table.rows:
        set_row_cant_split(row)


def shade_paragraph(paragraph: Any, fill: str) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_bottom_border(paragraph: Any, color: str = BORDER, size: str = "6") -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = p_bdr.find(qn("w:bottom"))
    if bottom is None:
        bottom = OxmlElement("w:bottom")
        p_bdr.append(bottom)
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), color)


def set_run_font(run: Any, size: float = 11, bold: bool | None = None, color: str = INK) -> None:
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold


def add_field(paragraph: Any, field: str) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = field
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, text, end])
    set_run_font(run, size=9, color=MUTED)


def configure_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for name, size, before, after, color in (
        ("Title", 24, 0, 8, INK),
        ("Heading 1", 16, 18, 10, BLUE),
        ("Heading 2", 13, 14, 7, BLUE),
        ("Heading 3", 12, 10, 5, INK),
    ):
        style = styles[name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def configure_section(section: Any) -> None:
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hp.paragraph_format.space_after = Pt(2)
    run = hp.add_run("中文叙事质量匿名排序问卷")
    set_run_font(run, size=9, bold=True, color=MUTED)
    set_bottom_border(hp, color="D7E0E6", size="4")

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fp.paragraph_format.space_before = Pt(2)
    run = fp.add_run("第 ")
    set_run_font(run, size=9, color=MUTED)
    add_field(fp, " PAGE ")
    run = fp.add_run(" 页")
    set_run_font(run, size=9, color=MUTED)


def paragraph(doc: Document, text: str = "", *, size: float = 11, bold: bool = False,
              color: str = INK, after: float = 6, align: Any = None) -> Any:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.25
    if align is not None:
        p.alignment = align
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)
    return p


def add_response_table(doc: Document, question_no: int) -> None:
    table = doc.add_table(rows=4, cols=2)
    table.style = "Table Grid"
    set_table_geometry(table, [1800, 7560])
    labels = ("综合排序", "判断信心", "简短理由", "记录编码")
    values = (
        "第1名：____　　第2名：____　　第3名：____　　第4名：____",
        "□ 1 很不确定　□ 2　□ 3　□ 4　□ 5 很确定",
        "________________________________________________________________",
        f"Q{question_no:02d}",
    )
    for index, (label, value) in enumerate(zip(labels, values)):
        left, right = table.rows[index].cells
        left.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        right.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        if index in (0, 3):
            for cell in (left, right):
                tc_pr = cell._tc.get_or_add_tcPr()
                shd = OxmlElement("w:shd")
                shd.set(qn("w:fill"), LIGHT)
                tc_pr.append(shd)
        lp = left.paragraphs[0]
        lp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        lp.paragraph_format.space_after = Pt(0)
        set_run_font(lp.add_run(label), size=10, bold=True, color=BLUE)
        rp = right.paragraphs[0]
        rp.paragraph_format.space_after = Pt(0)
        set_run_font(rp.add_run(value), size=10, color=INK)
    keep_table_rows_together(table)
    paragraph(doc, "请确认四个字母各使用一次，再进入下一题。", size=9, color=MUTED, after=0)


def shuffled_items(items: list[dict[str, Any]], prompt_id: str, seed: int) -> list[dict[str, Any]]:
    digest = hashlib.sha256(f"{seed}:{prompt_id}".encode("utf-8")).digest()
    rng = random.Random(int.from_bytes(digest[:8], "big"))
    result = list(items)
    rng.shuffle(result)
    return result


def build_document(work_dir: Path, output: Path, codebook_path: Path, seed: int) -> None:
    prompts = read_jsonl(work_dir / "prompts.jsonl")
    items = read_jsonl(work_dir / "blind_items.jsonl")
    grouped: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for item in items:
        grouped[item["prompt_id"]].append(item)
    if len(prompts) != 20 or len(items) != 80:
        raise ValueError("Expected exactly 20 prompts and 80 blinded items")
    if any(len(grouped[p["prompt_id"]]) != 4 for p in prompts):
        raise ValueError("Every prompt must have exactly four blinded items")

    doc = Document()
    configure_styles(doc)
    configure_section(doc.sections[0])
    props = doc.core_properties
    props.title = "中文叙事质量人工盲评排序问卷（20题）"
    props.subject = "来源隐藏的四篇文本排序问卷"
    props.author = "Narration Evaluation Study"
    props.last_modified_by = "Narration Evaluation Study"
    props.comments = "NV-PILOT20-RANK-v1"

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(54)
    title.paragraph_format.space_after = Pt(8)
    run = title.add_run("中文叙事质量人工盲评排序问卷")
    set_run_font(run, size=24, bold=True, color=INK)
    paragraph(doc, "20题 · 每题4篇匿名文本 · 版本 NV-PILOT20-RANK-v1", size=11,
              color=MUTED, after=30, align=WD_ALIGN_PARAGRAPH.CENTER)

    info = doc.add_table(rows=2, cols=2)
    info.style = "Table Grid"
    set_table_geometry(info, [4680, 4680])
    for cell in info.rows[0].cells + info.rows[1].cells:
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    fields = (("标注者编号：________________", "日期：____年__月__日"),
              ("开始时间：____:____", "结束时间：____:____"))
    for row, values in zip(info.rows, fields):
        for cell, value in zip(row.cells, values):
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            set_run_font(p.add_run(value), size=10.5, color=INK)

    doc.add_heading("填写说明", level=1)
    instructions = (
        "每题包含一个写作任务和4篇匿名候选文本。请完整阅读后，按照总体叙事质量从最好到最差排序。",
        "主要依据叙事完整性、情节与因果连贯、人物与视角、语言与节奏、情感或意象效果、原创性及任务遵循；不要只因篇幅或辞藻华丽程度加分。",
        "四篇文本必须分别占据第1至第4名，不得重复或空缺。请独立判断，不讨论、不搜索，也不要猜测作者身份。",
        "“判断信心”只反映你对本题排序的确定程度。简短理由可以留空，但出现难以区分的文本时建议填写。",
    )
    for text in instructions:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Inches(0.5)
        p.paragraph_format.first_line_indent = Inches(-0.25)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.25
        set_run_font(p.add_run(text), size=11, color=INK)

    note = paragraph(doc, "重要：文档中所有候选文本均已使用随机字母编号；排序结果中只填写 A、B、C、D。",
                     size=10.5, bold=True, color=BLUE, after=0)
    shade_paragraph(note, LIGHT)

    codebook: dict[str, Any] = {
        "instrument_version": "NV-PILOT20-RANK-v1",
        "seed": seed,
        "source_blinded": True,
        "questions": [],
    }
    answer_rows: list[dict[str, str]] = []

    for question_no, prompt in enumerate(prompts, start=1):
        doc.add_page_break()
        heading = doc.add_heading(f"第 {question_no:02d} 题", level=1)
        heading.paragraph_format.page_break_before = False
        meta = paragraph(doc, f"体裁：{prompt['genre']}　｜　请比较下列4篇候选文本",
                         size=9.5, bold=True, color=MUTED, after=7)
        meta.paragraph_format.keep_with_next = True
        task = paragraph(doc, "写作任务\n" + prompt["prompt"], size=10.5, color=INK, after=12)
        shade_paragraph(task, LIGHT)
        task.paragraph_format.keep_together = True

        ordered = shuffled_items(grouped[prompt["prompt_id"]], prompt["prompt_id"], seed)
        question_map: dict[str, Any] = {
            "question_no": question_no,
            "prompt_id": prompt["prompt_id"],
            "items": [],
        }
        for label, item in zip("ABCD", ordered):
            story_heading = doc.add_heading(f"候选文本 {label}", level=2)
            story_heading.paragraph_format.keep_with_next = True
            body = paragraph(doc, item["text"].strip(), size=11, color=INK, after=8)
            body.paragraph_format.line_spacing = 1.25
            body.paragraph_format.widow_control = True
            divider = paragraph(doc, "", after=3)
            set_bottom_border(divider, color="D7E0E6", size="4")
            question_map["items"].append({"label": label, "blind_id": item["blind_id"]})

        doc.add_heading("你的排序", level=2)
        add_response_table(doc, question_no)
        codebook["questions"].append(question_map)
        answer_rows.append({"question": f"Q{question_no:02d}"})

    doc.add_page_break()
    doc.add_heading("集中答题表", level=1)
    paragraph(doc, "如使用纸质版，可在此集中誊写每题答案。请与各题下方的排序记录保持一致。",
              size=10.5, color=MUTED, after=8)
    table = doc.add_table(rows=1 + len(answer_rows), cols=6)
    table.style = "Table Grid"
    set_table_geometry(table, [1200, 1600, 1600, 1600, 1600, 1760])
    headers = ("题号", "第1名", "第2名", "第3名", "第4名", "信心1—5")
    for cell, text in zip(table.rows[0].cells, headers):
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        tc_pr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), LIGHT)
        tc_pr.append(shd)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        set_run_font(p.add_run(text), size=9.5, bold=True, color=BLUE)
    set_repeat_table_header(table.rows[0])
    set_row_cant_split(table.rows[0])
    for row_index, record in enumerate(answer_rows, start=1):
        set_row_cant_split(table.rows[row_index])
        for column_index, cell in enumerate(table.rows[row_index].cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            value = record["question"] if column_index == 0 else ""
            set_run_font(p.add_run(value), size=10, bold=column_index == 0, color=INK)

    output.parent.mkdir(parents=True, exist_ok=True)
    codebook_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    codebook_path.write_text(json.dumps(codebook, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--work-dir", type=Path, default=DEFAULT_WORK_DIR)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--codebook", type=Path, default=DEFAULT_CODEBOOK)
    parser.add_argument("--seed", type=int, default=20260824)
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    build_document(args.work_dir.resolve(), args.output.resolve(), args.codebook.resolve(), args.seed)
    print(args.output.resolve())
    print(args.codebook.resolve())
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
