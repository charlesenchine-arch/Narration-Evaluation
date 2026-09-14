#!/usr/bin/env python3
"""Build the model-family aesthetic bias protocol as a polished Word document."""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs/FAMILY_AESTHETIC_BIAS_PROTOCOL.md"
OUTPUT = ROOT / "deliverables/大模型家族审美偏好与人类判断偏差实验协议_v4.docx"

FONT = "STHeiti"
MONO = "Courier New"
BLACK = "000000"
MUTED = "5E6872"
NAVY = "244A63"
PALE_BLUE = "EDF3F6"
PALE_GRAY = "F7F8F9"
GRID = "D9D9D9"


def set_run_font(run: Any, *, size: float = 11, bold: bool = False,
                 color: str = BLACK, font: str = FONT) -> None:
    run.font.name = font
    r_fonts = run._element.get_or_add_rPr().rFonts
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        r_fonts.set(qn(f"w:{key}"), font)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def set_cell_margins(cell: Any, top: int = 120, start: int = 140,
                     bottom: int = 120, end: int = 140) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_fill(cell: Any, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_borders(cell: Any, color: str = GRID, size: str = "6") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:color"), color)


def set_row_cant_split(row: Any) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:cantSplit")) is None:
        tr_pr.append(OxmlElement("w:cantSplit"))


def set_repeat_table_header(row: Any) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:tblHeader")) is None:
        tr_pr.append(OxmlElement("w:tblHeader"))


def add_page_field(paragraph: Any) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    value = OxmlElement("w:t")
    value.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, value, end])
    set_run_font(run, size=9, color=MUTED)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.85)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.92)
    section.right_margin = Inches(0.92)
    section.footer_distance = Inches(0.38)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(BLACK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    title = styles["Title"]
    title.font.name = FONT
    title._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    title.font.size = Pt(25)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string(BLACK)
    title.paragraph_format.space_after = Pt(12)
    title_p_pr = title._element.get_or_add_pPr()
    title_border = title_p_pr.find(qn("w:pBdr"))
    if title_border is not None:
        title_p_pr.remove(title_border)

    for name, size, before, after in (
        ("Heading 1", 16, 18, 8),
        ("Heading 2", 13, 14, 6),
        ("Heading 3", 11.5, 10, 4),
    ):
        style = styles[name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(BLACK)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.paragraph_format.space_before = Pt(0)
    set_run_font(footer.add_run("实验协议  "), size=9, color=MUTED)
    add_page_field(footer)


def add_inline(paragraph: Any, text: str, *, size: float = 11, color: str = BLACK) -> None:
    """Render a small subset of Markdown inline formatting."""
    pattern = re.compile(r"(\*\*.+?\*\*|`.+?`)")
    pos = 0
    for match in pattern.finditer(text):
        if match.start() > pos:
            set_run_font(paragraph.add_run(text[pos:match.start()]), size=size, color=color)
        token = match.group(0)
        if token.startswith("**"):
            set_run_font(paragraph.add_run(token[2:-2]), size=size, bold=True, color=color)
        else:
            set_run_font(paragraph.add_run(token[1:-1]), size=size - 0.3, color=color, font=MONO)
        pos = match.end()
    if pos < len(text):
        set_run_font(paragraph.add_run(text[pos:]), size=size, color=color)


def add_summary_table(doc: Document) -> None:
    heading = doc.add_heading("实验设计摘要", level=1)
    heading.paragraph_format.space_before = Pt(0)
    rows = [
        ("研究对象", "模型作为评审时的自身偏好 家族偏好 以及机器与人类判断偏差"),
        ("刺激材料", "同一写作任务下的匿名短篇文本 每篇 450 至 650 个中文非空字符 目标 500 字"),
        ("人类任务", "每人一道匿名两两比较 选择 A 更好 B 更好或难分高下"),
        ("模型任务", "每题六个 pair 每个 pair 换序复测 开放权重与闭源使用同一协议"),
        ("核心估计", "控制人类 pairwise 胜率后 各评审家族对各生成家族的额外偏好"),
        ("关键证据", "同家族跨模型效应 多模型方向一致 留一模型和题目分半后仍稳定"),
        ("文学锚点", "名著与授权人类文本独立检验机器和人类分歧 不进入核心家族偏好系数"),
    ]
    table = doc.add_table(rows=1 + len(rows), cols=2)
    table.autofit = False
    table.columns[0].width = Inches(1.4)
    table.columns[1].width = Inches(5.05)
    headers = ("项目", "冻结方案")
    for idx, (left, right) in enumerate([headers] + rows):
        row = table.rows[idx]
        set_row_cant_split(row)
        if idx == 0:
            set_repeat_table_header(row)
        for col, value in enumerate((left, right)):
            cell = row.cells[col]
            cell.width = Inches(1.4 if col == 0 else 5.05)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            set_cell_borders(cell)
            if idx == 0:
                set_cell_fill(cell, NAVY)
            elif idx % 2 == 0:
                set_cell_fill(cell, PALE_GRAY)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if col == 0 else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(0)
            set_run_font(p.add_run(value), size=10.2, bold=(idx == 0 or col == 0),
                         color="FFFFFF" if idx == 0 else BLACK)


def add_markdown_body(doc: Document, source: str) -> None:
    lines = source.splitlines()
    buffer: list[str] = []
    page_break_headings: set[str] = set()

    def flush() -> None:
        if not buffer:
            return
        text = " ".join(x.strip() for x in buffer).strip()
        buffer.clear()
        if not text:
            return
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.24
        p.paragraph_format.widow_control = True
        add_inline(p, text)

    skip_title = True
    for line in lines:
        if skip_title and line.startswith("# "):
            skip_title = False
            continue
        if line.startswith("更新日期："):
            continue
        if not line.strip():
            flush()
            continue
        heading = re.match(r"^(#{2,4})\s+(.+)$", line)
        if heading:
            flush()
            level = min(len(heading.group(1)) - 1, 3)
            title = heading.group(2).strip().replace("：", " ").replace(":", " ")
            p = doc.add_heading(title, level=level)
            if title in page_break_headings:
                p.paragraph_format.page_break_before = True
            continue
        bullet = re.match(r"^-\s+(.+)$", line)
        if bullet:
            flush()
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.left_indent = Inches(0.33)
            p.paragraph_format.first_line_indent = Inches(-0.2)
            p.paragraph_format.space_after = Pt(4)
            add_inline(p, bullet.group(1))
            continue
        numbered = re.match(r"^(\d+)\.\s+(.+)$", line)
        if numbered:
            flush()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.33)
            p.paragraph_format.first_line_indent = Inches(-0.2)
            p.paragraph_format.space_after = Pt(4)
            add_inline(p, f"{numbered.group(1)}.  {numbered.group(2)}")
            continue
        if line.startswith("`") and line.endswith("`"):
            flush()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.28)
            p.paragraph_format.right_indent = Inches(0.18)
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(8)
            add_inline(p, line)
            continue
        buffer.append(line)
    flush()


def build() -> None:
    source = SOURCE.read_text(encoding="utf-8")
    doc = Document()
    configure_document(doc)

    props = doc.core_properties
    props.title = "大模型家族审美偏好与人类判断偏差实验协议"
    props.subject = "中文短篇叙事的模型评审偏好实验设计"
    props.author = "Narration Evaluation Study"
    props.last_modified_by = "Narration Evaluation Study"
    props.comments = "Protocol version 4.0"

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(100)
    title_p_pr = title._p.get_or_add_pPr()
    title_border = title_p_pr.find(qn("w:pBdr"))
    if title_border is not None:
        title_p_pr.remove(title_border)
    set_run_font(title.add_run("大模型家族审美偏好\n与人类判断偏差实验协议"),
                 size=22, bold=True)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_before = Pt(10)
    subtitle.paragraph_format.space_after = Pt(18)
    set_run_font(subtitle.add_run("中文短篇叙事盲评研究"), size=14, color=MUTED)
    version = doc.add_paragraph()
    version.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(version.add_run("版本 4.0    2026 年 9 月 13 日"), size=10.5, color=MUTED)

    intro = doc.add_paragraph()
    intro.paragraph_format.space_before = Pt(78)
    intro.paragraph_format.left_indent = Inches(0.55)
    intro.paragraph_format.right_indent = Inches(0.55)
    intro.paragraph_format.line_spacing = 1.32
    set_run_font(intro.add_run(
        "本协议把人类盲评作为质量基准，检验模型评审是否在控制人类共识后仍偏好自己或同一模型家族的生成文本，并判断机器排序是否超出普通人类之间的正常分歧。"
    ), size=11.5)

    doc.add_page_break()
    add_summary_table(doc)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.line_spacing = 1.28
    set_run_font(p.add_run("主设计决定  "), size=11, bold=True)
    set_run_font(p.add_run(
        "人类和机器使用同一两两比较单位；人类每次只读两篇约 500 字文本，机器对同一 pair 位置互换复测。名著另设文学锚点实验，避免把模型记忆误当成家族审美。"
    ), size=11)

    add_markdown_body(doc, source)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
