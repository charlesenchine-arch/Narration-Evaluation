#!/usr/bin/env python3
"""Generate the Chinese pilot-20 feasibility/funding report and key-results chart."""

from __future__ import annotations

import argparse
import json
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


BLUE = "2E74B5"
DARK_BLUE = "163A5F"
LIGHT_BLUE = "EAF2F8"
LIGHT_GRAY = "F4F6F9"
MID_GRAY = "667085"
GOLD = "8A6500"
RED = "9B1C1C"
GREEN = "246B46"
FONT_LATIN = "Arial Unicode MS"
FONT_CJK = "Arial Unicode MS"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_run(run, *, size=11, bold=False, color="000000", italic=False) -> None:
    run.font.name = FONT_LATIN
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT_LATIN)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT_LATIN)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT_CJK)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def configure_style(style, *, size, bold=False, color="000000", before=0, after=6, line=1.25):
    style.font.name = FONT_LATIN
    style._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT_LATIN)
    style._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT_LATIN)
    style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT_CJK)
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor.from_string(color)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.line_spacing = line


def add_page_field(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    set_run(run, size=9, color=MID_GRAY)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)
    run = paragraph.add_run(" 页")
    set_run(run, size=9, color=MID_GRAY)


def set_table_geometry(table, widths: list[int], indent=120) -> None:
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    total = sum(widths)
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(total))
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_ind.set(qn("w:w"), str(indent))
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:type"), "dxa")
            tc_w.set(qn("w:w"), str(width))
            margins = tc_pr.find(qn("w:tcMar"))
            if margins is None:
                margins = OxmlElement("w:tcMar")
                tc_pr.append(margins)
            for side, value in (("top", 100), ("bottom", 100), ("start", 120), ("end", 120)):
                element = margins.find(qn(f"w:{side}"))
                if element is None:
                    element = OxmlElement(f"w:{side}")
                    margins.append(element)
                element.set(qn("w:w"), str(value))
                element.set(qn("w:type"), "dxa")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)


def fill_cell(cell, text: str, *, bold=False, color="000000", align=None, size=9.5):
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.15
    if align is not None:
        paragraph.alignment = align
    run = paragraph.add_run(text)
    set_run(run, size=size, bold=bold, color=color)


def add_heading(doc, text: str, level: int):
    paragraph = doc.add_heading(text, level=level)
    paragraph.paragraph_format.keep_with_next = True
    return paragraph


def add_body(doc, text: str, *, bold_lead: str | None = None):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if bold_lead and text.startswith(bold_lead):
        run = paragraph.add_run(bold_lead)
        set_run(run, bold=True, color=DARK_BLUE)
        run = paragraph.add_run(text[len(bold_lead) :])
        set_run(run)
    else:
        run = paragraph.add_run(text)
        set_run(run)
    return paragraph


def add_bullet(doc, text: str, *, level=0):
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.left_indent = Inches(0.375 + level * 0.25)
    paragraph.paragraph_format.first_line_indent = Inches(-0.194)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.208
    set_run(paragraph.add_run(text))
    return paragraph


def add_number(doc, text: str):
    paragraph = doc.add_paragraph(style="List Number")
    paragraph.paragraph_format.left_indent = Inches(0.375)
    paragraph.paragraph_format.first_line_indent = Inches(-0.194)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.208
    set_run(paragraph.add_run(text))
    return paragraph


def add_callout(doc, label: str, text: str, *, fill=LIGHT_BLUE, color=DARK_BLUE):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(10)
    paragraph.paragraph_format.left_indent = Inches(0.12)
    paragraph.paragraph_format.right_indent = Inches(0.12)
    paragraph.paragraph_format.line_spacing = 1.2
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)
    run = paragraph.add_run(f"{label}  ")
    set_run(run, size=11, bold=True, color=color)
    run = paragraph.add_run(text)
    set_run(run, size=11, color=color)
    return paragraph


def make_chart(data: dict, output: Path) -> None:
    font_path = Path("/Library/Fonts/NotoSansCJKsc-VF.ttf")
    if not font_path.exists():
        font_path = Path("/System/Library/Fonts/STHeiti Light.ttc")

    def font(size: int, bold: bool = False):
        try:
            return ImageFont.truetype(str(font_path), size=size, index=0)
        except Exception:
            return ImageFont.load_default()

    width, height = 2376, 968
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    title_font, panel_font, label_font, value_font = font(42, True), font(31, True), font(24), font(25, True)
    draw.text((width // 2, 38), "20题预实验：来源线索强，代理评审显著偏向当前机器样本", fill="#163A5F", font=title_font, anchor="ma")

    panels = [(100, 160, 1120, 840), (1256, 160, 2276, 840)]
    for left, top, right, bottom in panels:
        draw.rounded_rectangle((left, top, right, bottom), radius=18, fill="#FBFCFE", outline="#D0D5DD", width=2)

    probes = data["source_probe"]
    labels = ["字形TF-IDF", "仅长度", "表层特征"]
    auc = [
        probes["char_tfidf_probe"]["roc_auc"],
        probes["length_only_probe"]["roc_auc"],
        probes["surface_feature_probe"]["roc_auc"],
    ]
    draw.text((610, 202), "来源可识别性（ROC AUC）", fill="#101828", font=panel_font, anchor="ma")
    chart_left, chart_top, chart_right, chart_bottom = 190, 286, 1050, 730
    draw.line((chart_left, chart_bottom, chart_right, chart_bottom), fill="#667085", width=2)
    draw.line((chart_left, chart_top, chart_left, chart_bottom), fill="#667085", width=2)
    for tick in [0.5, 0.75, 1.0]:
        y = chart_bottom - (tick - 0.45) / 0.57 * (chart_bottom - chart_top)
        draw.line((chart_left, y, chart_right, y), fill="#E4E7EC", width=2)
        draw.text((chart_left - 18, y), f"{tick:.2f}", fill="#667085", font=label_font, anchor="rm")
    colors = ["#2E74B5", "#8A6500", "#7C8DA6"]
    centers = [350, 610, 870]
    for center, label, value, color in zip(centers, labels, auc, colors):
        y = chart_bottom - (value - 0.45) / 0.57 * (chart_bottom - chart_top)
        draw.rounded_rectangle((center - 72, y, center + 72, chart_bottom), radius=8, fill=color)
        draw.text((center, y - 16), f"{value:.3f}", fill="#101828", font=value_font, anchor="ms")
        draw.text((center, chart_bottom + 28), label, fill="#344054", font=label_font, anchor="ma")

    reward = data["reward_model"]["pairwise"]["overall"]
    qwen = data["qwen_judge"]["pairwise"]["overall"]
    averaged = data["qwen_judge"]["reversed_position_audit"]["position_averaged_score_preference"]
    summaries = [reward, qwen, averaged]
    labels = ["Reward 60组", "Qwen原始 60组", "Qwen双位置 20组"]
    values = [summary["human_win_equivalent"] for summary in summaries]
    draw.text((1766, 202), "代理模型中的人类文本胜率", fill="#101828", font=panel_font, anchor="ma")
    chart_left, chart_top, chart_right, chart_bottom = 1346, 286, 2206, 730
    draw.line((chart_left, chart_bottom, chart_right, chart_bottom), fill="#667085", width=2)
    draw.line((chart_left, chart_top, chart_left, chart_bottom), fill="#667085", width=2)
    for tick in [0.0, 0.25, 0.5]:
        y = chart_bottom - tick / 0.62 * (chart_bottom - chart_top)
        draw.line((chart_left, y, chart_right, y), fill="#E4E7EC", width=2)
        draw.text((chart_left - 18, y), f"{tick:.2f}", fill="#667085", font=label_font, anchor="rm")
    baseline_y = chart_bottom - 0.5 / 0.62 * (chart_bottom - chart_top)
    for x in range(chart_left, chart_right, 24):
        draw.line((x, baseline_y, min(x + 12, chart_right), baseline_y), fill="#98A2B3", width=3)
    colors = ["#2E74B5", "#667085", "#246B46"]
    centers = [1506, 1766, 2026]
    for center, label, value, summary, color in zip(centers, labels, values, summaries, colors):
        y = chart_bottom - value / 0.62 * (chart_bottom - chart_top)
        draw.rounded_rectangle((center - 72, y, center + 72, chart_bottom), radius=8, fill=color)
        low, high = summary["human_win_equivalent_95ci"]
        low_y = chart_bottom - low / 0.62 * (chart_bottom - chart_top)
        high_y = chart_bottom - high / 0.62 * (chart_bottom - chart_top)
        draw.line((center, high_y, center, low_y), fill="#101828", width=4)
        draw.line((center - 18, high_y, center + 18, high_y), fill="#101828", width=4)
        draw.line((center - 18, low_y, center + 18, low_y), fill="#101828", width=4)
        draw.text((center, y - 18), f"{value:.1%}", fill="#101828", font=value_font, anchor="ms")
        draw.text((center, chart_bottom + 28), label, fill="#344054", font=font(21), anchor="ma")
    draw.text((2175, baseline_y - 12), "无偏好 0.5", fill="#667085", font=font(20), anchor="rs")
    output.parent.mkdir(parents=True, exist_ok=True)
    image.save(output, dpi=(220, 220))


def build_doc(data: dict, chart: Path, output: Path) -> None:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    configure_style(doc.styles["Normal"], size=11, after=6, line=1.25)
    configure_style(doc.styles["Heading 1"], size=16, bold=True, color=BLUE, before=16, after=8, line=1.1)
    configure_style(doc.styles["Heading 2"], size=13, bold=True, color=BLUE, before=12, after=6, line=1.1)
    configure_style(doc.styles["Heading 3"], size=12, bold=True, color=DARK_BLUE, before=8, after=4, line=1.1)
    configure_style(doc.styles["List Bullet"], size=11, after=4, line=1.208)
    configure_style(doc.styles["List Number"], size=11, after=4, line=1.208)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_run(header.add_run("NARRATION EVALUATION  |  PILOT 20"), size=9, bold=True, color=MID_GRAY)
    add_page_field(section.footer.paragraphs[0])

    # proposal_centerpiece first-page header pattern
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(36)
    p.paragraph_format.space_after = Pt(10)
    set_run(p.add_run("研究预实验与资助建议"), size=12, bold=True, color=GOLD)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    set_run(p.add_run("中文叙事质量评估中的来源偏差"), size=26, bold=True, color=DARK_BLUE)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(26)
    set_run(p.add_run("20题盲测可行性报告"), size=16, color=BLUE)

    meta = doc.add_table(rows=2, cols=2)
    values = [
        ("实验规模", "20题 · 80篇 · 60组人机配对"),
        ("本地基线", "Llama-3 8B Reward + Qwen3-VL 8B"),
        ("计算成本", "服务器本地推理；API费用为0"),
        ("报告日期", "2026年8月21日"),
    ]
    for cell, (label, value) in zip([c for row in meta.rows for c in row.cells], values):
        set_cell_shading(cell, LIGHT_GRAY)
        paragraph = cell.paragraphs[0]
        paragraph.paragraph_format.space_after = Pt(1)
        set_run(paragraph.add_run(label + "\n"), size=9, bold=True, color=MID_GRAY)
        set_run(paragraph.add_run(value), size=10.5, bold=True, color=DARK_BLUE)
    set_table_geometry(meta, [4680, 4680], indent=120)

    doc.add_paragraph().paragraph_format.space_after = Pt(10)
    add_callout(
        doc,
        "决策建议",
        "建议进入“有条件资助”的第二阶段，但资金应优先用于数据校准与人工盲评，而不是立即训练大规模评分器。技术路径已跑通，研究问题也真实存在；当前20题结果尚不能证明来源偏差。",
        fill=LIGHT_BLUE,
        color=DARK_BLUE,
    )

    add_heading(doc, "一、这次预实验回答了什么", 1)
    add_body(
        doc,
        "核心问题不是“模型能否给一个分数”，而是：当人类文本与机器文本质量相当时，评估器是否仍会根据来源痕迹产生系统性偏好。等级、分数、维度和评语只是解释判断的输出层；论文主线是质量与来源的解耦。",
    )
    add_body(
        doc,
        "本轮20题实验先验证两个前提：第一，现有流水线能否无API费用地完成盲配、评分和稳健性检查；第二，当前文本池是否已经具备直接研究来源偏差的条件。答案分别是“能”和“还不具备”。",
    )

    add_heading(doc, "二、实验设计", 1)
    for text in [
        "20个反向构造的中文叙事题目，每题对应1篇现有人类文本和3篇机器文本。",
        "机器条件包括Gemini 3.5 Flash Lite、Gemini 3.6 Flash与DeepSeek聊天端（关闭深度思考和搜索）。",
        "共80篇文本、60组同题人机配对；左右位置在主配对中总体平衡，评审时不提供来源或生成器标签。",
        "本地Llama-3 8B reward model为80篇文本打分；本地Qwen3-VL 8B完成60组中文编辑式盲评。",
        "另外抽取每题1组、共20组反转A/B位置重评，用于测量位置偏差。",
    ]:
        add_bullet(doc, text)

    add_heading(doc, "三、关键结果", 1)
    table = doc.add_table(rows=1, cols=3)
    headers = ["观察", "结果", "应如何解释"]
    for cell, text in zip(table.rows[0].cells, headers):
        set_cell_shading(cell, LIGHT_GRAY)
        fill_cell(cell, text, bold=True, color=DARK_BLUE, align=WD_ALIGN_PARAGRAPH.CENTER, size=9.5)
    rows = [
        ("来源可识别", "字形模型 AUC 0.959；仅长度 AUC 0.963", "人机来源痕迹非常强，且长度几乎单独解释来源"),
        ("Reward偏好", "机器胜 59/60；模型间分差约 +4.95（机器）", "当前机器样本明显更符合通用偏好模型"),
        ("Qwen盲评", "机器胜 53/60；人类胜率 11.7%", "中文编辑式代理评审同样强烈偏向机器样本"),
        ("跨模型一致", "方向一致率 86.7%", "结果不只是单一模型的偶然输出"),
        ("位置稳健性", "反转后方向一致 70%；显示A被选 15/20", "Qwen存在明显位置敏感，正式实验必须双位置汇总"),
        ("双位置汇总", "机器胜 19/20；人类胜率 5.0%", "消除部分位置影响后，当前样本差距仍很大"),
    ]
    for observation, result, meaning in rows:
        cells = table.add_row().cells
        fill_cell(cells[0], observation, bold=True, color=DARK_BLUE)
        fill_cell(cells[1], result, color="000000")
        fill_cell(cells[2], meaning, color="000000")
    set_repeat_table_header(table.rows[0])
    set_table_geometry(table, [1900, 2700, 4760], indent=120)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    set_run(p.add_run("图1　来源可识别性与代理模型偏好"), size=9.5, bold=True, color=MID_GRAY)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(chart), width=Inches(6.35))
    p.paragraph_format.space_after = Pt(2)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(p.add_run("注：误差线为按题目分组自助法95%区间；代理模型结果不是人类金标准。"), size=8.5, color=MID_GRAY)

    add_heading(doc, "四、最重要的科学结论", 1)
    add_callout(
        doc,
        "不是“发现机器偏好”",
        "本轮只能说明：当前机器文本在两个代理评审中显著高于当前人类文本。由于人类文本尚未经过独立质量审核，而且机器文本普遍更长，因此无法判断评审器是在识别真实质量差异，还是在利用来源线索。",
        fill="FFF4E5",
        color=GOLD,
    )
    add_body(
        doc,
        "这不是预实验失败，反而准确定位了论文最难、也最有价值的部分：必须先建立“人类高质量”和“机器高质量”可比的质量重叠区，才能估计控制质量后的来源残差。若跳过这一步直接训练评分器，论文很容易退化为“区分当前数据集的人和机器”。",
    )
    add_body(
        doc,
        "长度相关性也不能简单等同于质量。80篇文本中，reward分数与长度的Spearman相关为0.694，Qwen平均分与长度的相关为0.421；但探索性线性校正后，机器优势仍然存在。这说明长度是显著混杂因素，却不是唯一解释。真正的判定仍需长度匹配生成和独立人类盲评。",
    )
    add_body(
        doc,
        "此外，Qwen有21/60条输出出现“偏好字段与分数/评语不一致”，本报告按分数重新计算偏好并保留原始记录。该现象与位置敏感共同说明：LLM-as-a-Judge只能作为基线，不能替代人类标签。",
    )

    add_heading(doc, "五、为什么值得申请下一阶段资助", 1)
    for text in [
        "可执行：盲化、同题配对、生成器留出、来源探针、reward基线、LLM裁判和位置反转检查均已实现。",
        "可证伪：若质量匹配后来源效应消失，论文应诚实报告“没有来源偏差”；若仍存在，则得到核心发现。",
        "有现实价值：现有自动评审可能把长度、模板化完整度或生成器风格误当作文学质量，直接影响模型训练与内容筛选。",
        "成本可控：服务器已有H100与开源模型；下一阶段主要支出应是人工标注、文本授权与数据清理。",
    ]:
        add_bullet(doc, text)

    add_heading(doc, "六、建议资助的最小下一阶段", 1)
    add_number(doc, "先修复这20题：逐篇审核人类文本的完整性、版权/许可、题目一致性和明显噪声；不合格文本替换，不扩规模。")
    add_number(doc, "按每篇人类文本长度的±10%重新生成机器文本，并固定无搜索、无外部工具和统一采样协议；保留至少3个生成器条件。")
    add_number(doc, "对60组同题人机配对做双盲人工评审，每组至少5名评审者，共约300份成对判断；先测评审一致性与质量重叠。")
    add_number(doc, "模型评审一律双位置运行并汇总；reward、开源LLM judge和简单长度/表层特征模型共同作为基线。")
    add_number(doc, "达到继续标准后再扩到40—100题；未达到则先修数据，不进入2000篇规模，也不训练最终评分器。")

    add_heading(doc, "七、第二阶段的继续/停止标准", 1)
    gate = doc.add_table(rows=1, cols=3)
    for cell, text in zip(gate.rows[0].cells, ["检查点", "继续条件", "不满足时的动作"]):
        set_cell_shading(cell, LIGHT_GRAY)
        fill_cell(cell, text, bold=True, color=DARK_BLUE, align=WD_ALIGN_PARAGRAPH.CENTER)
    gates = [
        ("人类数据质量", "审计通过且无明显截断/任务不一致", "替换文本，不扩数据"),
        ("长度控制", "机器文本相对人类文本控制在±10%", "重新生成对应样本"),
        ("人工可靠性", "成对偏好达到可解释的一致性", "优化指南、培训评审者"),
        ("质量重叠", "存在足量高机低人、低机高人及相近质量配对", "补采边界样本"),
        ("论文信号", "控制人类质量后仍有来源残差或未见生成器崩溃", "转向评测协议/负结果论文"),
    ]
    for checkpoint, go, fallback in gates:
        cells = gate.add_row().cells
        fill_cell(cells[0], checkpoint, bold=True, color=DARK_BLUE)
        fill_cell(cells[1], go)
        fill_cell(cells[2], fallback)
    set_repeat_table_header(gate.rows[0])
    set_table_geometry(gate, [1900, 4400, 3060], indent=120)

    add_heading(doc, "八、最终判断", 1)
    add_body(
        doc,
        "这20题已经足以证明项目值得做一个受控的小规模人工阶段，但不足以支持论文结论，更不足以直接训练最终评估器。最合理的资助叙事不是“我们已经发现模型偏见”，而是“我们已经用零API成本跑通方法，并发现现有数据会把来源、长度与真实质量严重混在一起；下一阶段将用人工金标准和长度匹配实验解决这个关键识别问题”。",
    )
    add_callout(
        doc,
        "资助定位",
        "建议申请“方法验证与数据校准型”小额启动经费。主要预算：人工盲评、数据授权/清理、长度匹配再生成；算力预算可以很低。",
        fill="EAF6EF",
        color=GREEN,
    )

    add_heading(doc, "附录：复现实验资产", 1)
    assets = [
        "20题盲测清单与私有来源键（80篇文本）",
        "60组人机盲配与120组同题全配对",
        "来源探针、reward评分、Qwen盲评与20组反转审计原始结果",
        "分组自助法置信区间和长度探索性校正脚本",
        "本地服务器开源模型运行脚本；无需付费API",
    ]
    for asset in assets:
        add_bullet(doc, asset)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    set_run(p.add_run("限制声明："), size=9, bold=True, color=RED)
    set_run(
        p.add_run("人类样本来自现有数据混池，尚处于暂定审核状态；本报告中的模型判断均为代理结果，不得写作人类质量结论或因果偏差结论。"),
        size=9,
        color=RED,
    )

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--analysis",
        type=Path,
        default=Path("outputs/human_pilot_20260821/pilot20_combined_analysis.json"),
    )
    parser.add_argument(
        "--chart",
        type=Path,
        default=Path("outputs/human_pilot_20260821/pilot20_key_results.png"),
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=Path("deliverables/中文叙事来源偏差_20题预实验与资助建议.docx"),
    )
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    data = json.loads(args.analysis.read_text(encoding="utf-8"))
    make_chart(data, args.chart)
    build_doc(data, args.chart, args.output)
    print(args.chart.resolve())
    print(args.output.resolve())
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
